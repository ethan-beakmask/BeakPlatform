"""
Spec Formulate - 匯出寫入器

將 spec_formulate 的 fields JSONB 格式匯出為 Excel、CSV、DOCX。
適配自 sql_multifaceted 的 writers，配合 BeakPlatform 的欄位結構。
"""
import csv
import os
import logging
import tempfile
from datetime import datetime

logger = logging.getLogger(__name__)


# ============================================================
# Excel Writer
# ============================================================

def export_excel(fields, spec_name, output_path=None):
    """
    匯出為 Excel 規格書

    結構:
    - Sheet1: 欄位規格表（Label, Field Key, Type, PG Type, PII, 必填, 說明）
    - 第1列: 表頭（粗體、淺灰底）
    - 第2列起: 資料列

    Args:
        fields: spec fields list (JSONB 格式)
        spec_name: 規格名稱（用於 sheet name）
        output_path: 輸出路徑，None 時自動建立暫存檔

    Returns:
        str: 輸出檔案路徑
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    if output_path is None:
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        output_path = tmp.name

    wb = Workbook()
    ws = wb.active
    ws.title = (spec_name or 'spec')[:31]

    # 樣式
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    mono_font = Font(name='Consolas', size=10)
    pii_fill = PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    headers = ['#', 'Label', 'Field Key', 'Type', 'PG Type', 'PII', '必填', '說明']
    col_widths = [5, 18, 18, 14, 16, 6, 6, 30]

    # 表頭
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = col_widths[i - 1]

    # 資料列
    for row_idx, f in enumerate(fields, start=2):
        constraints = f.get('constraints') or {}
        values = [
            row_idx - 1,
            f.get('label', ''),
            f.get('field_key', ''),
            _formio_type_label(f.get('formio_type', '')),
            f.get('pg_type', ''),
            'Y' if f.get('is_pii') else '',
            'Y' if constraints.get('required') else '',
            f.get('description', ''),
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            # 等寬字體給 key 和 pg_type
            if col_idx in (3, 5):
                cell.font = mono_font
            # 序號置中
            if col_idx in (1, 6, 7):
                cell.alignment = Alignment(horizontal='center')
            # PII 列底色
            if f.get('is_pii'):
                cell.fill = pii_fill

    # 凍結表頭
    ws.freeze_panes = 'A2'

    wb.save(output_path)
    return output_path


# ============================================================
# CSV Writer
# ============================================================

def export_csv(fields, spec_name, output_path=None, encoding='utf-8-sig'):
    """
    匯出為 CSV

    Args:
        fields: spec fields list
        spec_name: 規格名稱（未使用，保持介面一致）
        output_path: 輸出路徑
        encoding: 編碼，預設 UTF-8 with BOM

    Returns:
        str: 輸出檔案路徑
    """
    if output_path is None:
        tmp = tempfile.NamedTemporaryFile(suffix='.csv', delete=False)
        tmp.close()
        output_path = tmp.name

    headers = ['#', 'Label', 'Field Key', 'Type', 'PG Type', 'PII', '必填', '說明']

    with open(output_path, 'w', encoding=encoding, newline='') as fp:
        writer = csv.writer(fp)
        writer.writerow(headers)

        for idx, f in enumerate(fields, start=1):
            constraints = f.get('constraints') or {}
            writer.writerow([
                idx,
                f.get('label', ''),
                f.get('field_key', ''),
                _formio_type_label(f.get('formio_type', '')),
                f.get('pg_type', ''),
                'Y' if f.get('is_pii') else '',
                'Y' if constraints.get('required') else '',
                f.get('description', ''),
            ])

    return output_path


# ============================================================
# DOCX Writer
# ============================================================

# 預設前後文
DEFAULT_PREAMBLE = (
    '本章節定義系統所使用之資料表結構規格。各資料表之欄位名稱、'
    '資料型別、長度限制、是否允許空值、主鍵設定及欄位用途說明如下。'
)

DEFAULT_POSTSCRIPT = (
    '如需異動資料表結構，應先更新規格書，經審核確認後再進行資料庫變更。'
)

# 樣式常數
_FONT_EN = 'Calibri'
_FONT_ZH = 'Noto Sans CJK TC'
_COLOR_HEADER_BG = 'D6E4F0'
_COLOR_PII_BG = 'FCE4EC'
_COLOR_BORDER = '8EAADB'


def export_docx(fields, spec_name, output_path=None,
                preamble=None, postscript=None):
    """
    匯出為 Word 規格文件

    Args:
        fields: spec fields list
        spec_name: 規格名稱
        output_path: 輸出路徑
        preamble: 前文（None 使用預設）
        postscript: 後文（None 使用預設）

    Returns:
        str: 輸出檔案路徑
    """
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn

    if output_path is None:
        tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        tmp.close()
        output_path = tmp.name

    preamble = preamble if preamble is not None else DEFAULT_PREAMBLE
    postscript = postscript if postscript is not None else DEFAULT_POSTSCRIPT

    doc = Document()

    # 預設字型
    style = doc.styles['Normal']
    style.font.name = _FONT_EN
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor.from_string('333333')
    style.element.rPr.rFonts.set(qn('w:eastAsia'), _FONT_ZH)

    # 頁面邊距
    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)

    # 標題
    _add_heading(doc, spec_name or '資料表規格書', level=1)

    # 前文
    if preamble.strip():
        for para_text in preamble.strip().split('\n\n'):
            p = doc.add_paragraph(para_text.strip())
            _set_para_font(p)
            p.paragraph_format.space_after = Pt(6)

    doc.add_paragraph()

    # 欄位表格
    table_headers = ['#', 'Label', 'Field Key', 'Type', 'PG Type',
                     'PII', '必填', '說明']
    col_count = len(table_headers)
    row_count = len(fields) + 1

    table = doc.add_table(rows=row_count, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    widths = [Cm(0.8), Cm(2.8), Cm(2.8), Cm(2.2), Cm(2.5),
              Cm(0.8), Cm(0.8), Cm(4.3)]

    # 表頭
    header_row = table.rows[0]
    for i, ht in enumerate(table_headers):
        cell = header_row.cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(ht)
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = _FONT_EN
        run.font.color.rgb = RGBColor.from_string('1F3864')
        run.element.rPr.rFonts.set(qn('w:eastAsia'), _FONT_ZH)
        _set_cell_bg(cell, _COLOR_HEADER_BG, qn)
        cell.width = widths[i]

    # 資料列
    for row_idx, f in enumerate(fields):
        row = table.rows[row_idx + 1]
        constraints = f.get('constraints') or {}
        values = [
            str(row_idx + 1),
            f.get('label', ''),
            f.get('field_key', ''),
            _formio_type_label(f.get('formio_type', '')),
            f.get('pg_type', ''),
            'Y' if f.get('is_pii') else '',
            'Y' if constraints.get('required') else '',
            f.get('description', ''),
        ]

        for col_idx, val in enumerate(values):
            cell = row.cells[col_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            run.font.name = _FONT_EN
            run.element.rPr.rFonts.set(qn('w:eastAsia'), _FONT_ZH)
            if col_idx in (0, 5, 6):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.width = widths[col_idx]

        # PII 列底色
        if f.get('is_pii'):
            for cell in row.cells:
                _set_cell_bg(cell, _COLOR_PII_BG, qn)

    # 表格框線
    _set_table_borders(table, qn)

    # 後文
    if postscript.strip():
        doc.add_paragraph()
        for para_text in postscript.strip().split('\n\n'):
            p = doc.add_paragraph(para_text.strip())
            _set_para_font(p)
            p.paragraph_format.space_after = Pt(6)

    doc.save(output_path)
    return output_path


# ============================================================
# 內部工具
# ============================================================

_FORMIO_TYPE_LABELS = {
    'textfield': '文字',
    'textarea': '多行文字',
    'number': '數字',
    'checkbox': '核取方塊',
    'day': '日期',
    'datetime': '日期時間',
    'email': '電子郵件',
    'phoneNumber': '電話',
    'select': '下拉選單',
    'radio': '單選按鈕',
    'selectboxes': '複選框',
    'file': '檔案',
    'signature': '簽名',
    'hidden': '隱藏',
    'currency': '貨幣',
    'url': 'URL',
    'tags': '標籤',
    'datagrid': '資料表格',
    'editgrid': '編輯表格',
}


def _formio_type_label(ftype):
    """取得 FormIO 型別的中文標籤"""
    return _FORMIO_TYPE_LABELS.get(ftype, ftype)


def _add_heading(doc, text, level=1):
    """新增標題"""
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = _FONT_EN
    run.font.color.rgb = RGBColor.from_string('1F3864')
    run.element.rPr.rFonts.set(qn('w:eastAsia'), _FONT_ZH)


def _set_para_font(paragraph):
    """設定段落字型"""
    from docx.oxml.ns import qn
    for run in paragraph.runs:
        run.font.name = _FONT_EN
        run.element.rPr.rFonts.set(qn('w:eastAsia'), _FONT_ZH)


def _set_cell_bg(cell, color_hex, qn_func):
    """設定儲存格背景色"""
    shading = cell._element.get_or_add_tcPr()
    elem = shading.makeelement(qn_func('w:shd'), {
        qn_func('w:val'): 'clear',
        qn_func('w:color'): 'auto',
        qn_func('w:fill'): color_hex,
    })
    shading.append(elem)


def _set_table_borders(table, qn_func):
    """設定表格框線"""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.makeelement(
        qn_func('w:tblPr'), {})

    borders = tblPr.makeelement(qn_func('w:tblBorders'), {})
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders.makeelement(qn_func(f'w:{edge}'), {
            qn_func('w:val'): 'single',
            qn_func('w:sz'): '4',
            qn_func('w:space'): '0',
            qn_func('w:color'): _COLOR_BORDER,
        })
        borders.append(element)

    for existing in tblPr.findall(qn_func('w:tblBorders')):
        tblPr.remove(existing)

    tblPr.append(borders)
