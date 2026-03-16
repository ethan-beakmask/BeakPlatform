"""
Spec Formulate - 格式讀取器

讀取 Excel/CSV 格式的欄位定義，轉換為 SPEC fields 格式。
與 export_writers.py 互為逆向操作，確保雙向轉換。

讀取的欄位格式（與匯出一致）:
# | Label | Field Key | Type | PG Type | PII | 必填 | 說明
"""
import csv
import io
import logging

logger = logging.getLogger(__name__)


# ============================================================
# FormIO 型別對應（與 export_writers._FORMIO_TYPE_LABELS 互逆）
# ============================================================

_LABEL_TO_FORMIO_TYPE = {
    '文字': 'textfield',
    '多行文字': 'textarea',
    '數字': 'number',
    '核取方塊': 'checkbox',
    '日期': 'day',
    '日期時間': 'datetime',
    '電子郵件': 'email',
    '電話': 'phoneNumber',
    '下拉選單': 'select',
    '單選按鈕': 'radio',
    '複選框': 'selectboxes',
    '檔案': 'file',
    '簽名': 'signature',
    '隱藏': 'hidden',
    '貨幣': 'currency',
    'url': 'url',
    '標籤': 'tags',
    '資料表格': 'datagrid',
    '編輯表格': 'editgrid',
}

# 合法的 FormIO 型別英文值（用於直接比對）
_VALID_FORMIO_TYPES = {
    'textfield', 'textarea', 'number', 'checkbox', 'day', 'datetime',
    'email', 'phoneNumber', 'select', 'radio', 'selectboxes',
    'file', 'signature', 'hidden', 'currency', 'url', 'tags',
    'datagrid', 'editgrid',
}

# 英文值小寫→正確大小寫對應（處理 phoneNumber 等特殊大小寫）
_FORMIO_TYPE_NORMALIZE = {t.lower(): t for t in _VALID_FORMIO_TYPES}

# FormIO Type → PG Type 預設對應
_PG_DEFAULTS = {
    'textfield': 'VARCHAR(500)',
    'textarea': 'TEXT',
    'number': 'NUMERIC',
    'checkbox': 'BOOLEAN',
    'day': 'DATE',
    'datetime': 'TIMESTAMP',
    'email': 'VARCHAR(200)',
    'phoneNumber': 'VARCHAR(50)',
    'select': 'VARCHAR(500)',
    'radio': 'VARCHAR(200)',
    'selectboxes': 'JSONB',
    'file': 'JSONB',
    'signature': 'TEXT',
    'hidden': 'TEXT',
    'currency': 'NUMERIC(15,2)',
    'url': 'VARCHAR(1000)',
    'tags': 'JSONB',
    'datagrid': 'JSONB',
    'editgrid': 'JSONB',
}

# 表頭別名映射（支援中英文表頭）
_HEADER_ALIASES = {
    'label': ['label', '標籤', '顯示名稱', '名稱'],
    'field_key': ['field key', 'field_key', 'fieldkey', 'key', '欄位名', '欄位key'],
    'type': ['type', '型別', '類型', 'formio type', 'formio_type'],
    'pg_type': ['pg type', 'pg_type', 'pgtype', 'sql type', 'sql_type', '資料庫型別'],
    'pii': ['pii', 'is_pii', '個資', '個人識別'],
    'required': ['必填', 'required', '必要'],
    'description': ['說明', 'description', 'desc', '描述', '備註'],
}


# ============================================================
# 公開 API
# ============================================================

def read_excel_to_spec_fields(file_storage):
    """
    讀取 Excel 檔案，轉換為 SPEC fields 格式。

    Args:
        file_storage: Flask FileStorage 或 file-like object

    Returns:
        list[dict]: SPEC fields 陣列

    Raises:
        ValueError: 格式無法辨識時
    """
    from openpyxl import load_workbook

    content = io.BytesIO(file_storage.read())
    wb = load_workbook(content, read_only=True, data_only=True)
    try:
        ws = wb.active

        headers = _read_excel_headers(ws)
        if not headers:
            raise ValueError('Excel 第一列無表頭資料')

        mapping = _find_column_mapping(headers)
        if mapping is None:
            raise ValueError(
                f'無法辨識表頭格式。找到的表頭: {headers[:10]}。'
                '需要至少包含 Label 或 Field Key 欄位。'
            )

        fields = []
        row = 2
        empty_streak = 0
        while row <= (ws.max_row or 1):
            row_data = []
            for c in range(1, len(headers) + 1):
                row_data.append(ws.cell(row=row, column=c).value)

            label_val = _safe_get(row_data, mapping.get('label'))
            fkey_val = _safe_get(row_data, mapping.get('field_key'))

            if not label_val and not fkey_val:
                empty_streak += 1
                if empty_streak >= 3:
                    break
                row += 1
                continue

            empty_streak = 0
            field = _row_to_spec_field(row_data, mapping, len(fields))
            if field:
                fields.append(field)

            row += 1

        if not fields:
            raise ValueError('Excel 中沒有有效的欄位資料')

        return fields

    finally:
        wb.close()


def read_csv_to_spec_fields(file_storage):
    """
    讀取 CSV 檔案，轉換為 SPEC fields 格式。
    支援 UTF-8 with BOM / UTF-8 編碼。

    Args:
        file_storage: Flask FileStorage 或 file-like object

    Returns:
        list[dict]: SPEC fields 陣列

    Raises:
        ValueError: 格式無法辨識時
    """
    raw = file_storage.read()

    if raw[:3] == b'\xef\xbb\xbf':
        text = raw.decode('utf-8-sig')
    else:
        try:
            text = raw.decode('utf-8')
        except UnicodeDecodeError:
            raise ValueError('CSV 編碼無法辨識，僅支援 UTF-8 或 UTF-8 with BOM')

    reader = csv.reader(io.StringIO(text))

    try:
        headers = next(reader)
    except StopIteration:
        raise ValueError('CSV 檔案是空的')

    headers = [h.strip() for h in headers]
    mapping = _find_column_mapping(headers)
    if mapping is None:
        raise ValueError(
            f'無法辨識表頭格式。找到的表頭: {headers[:10]}。'
            '需要至少包含 Label 或 Field Key 欄位。'
        )

    fields = []
    for row_data in reader:
        label_val = _safe_get(row_data, mapping.get('label'))
        fkey_val = _safe_get(row_data, mapping.get('field_key'))

        if not label_val and not fkey_val:
            continue

        field = _row_to_spec_field(row_data, mapping, len(fields))
        if field:
            fields.append(field)

    if not fields:
        raise ValueError('CSV 中沒有有效的欄位資料')

    return fields


# ============================================================
# 內部工具
# ============================================================

def _read_excel_headers(ws):
    """讀取 Excel 第一列表頭"""
    headers = []
    col = 1
    while col <= 50:
        val = ws.cell(row=1, column=col).value
        if val is None:
            if col > 1:
                next_val = ws.cell(row=1, column=col + 1).value
                if next_val is None:
                    break
            else:
                break
        headers.append(str(val).strip() if val is not None else '')
        col += 1
    return headers


def _find_column_mapping(headers):
    """
    分析表頭，建立欄位名到索引的映射。

    Returns:
        dict 或 None（無法辨識時）
    """
    mapping = {}
    normalized = [str(h).strip().lower() for h in headers]

    for field, alias_list in _HEADER_ALIASES.items():
        for i, h in enumerate(normalized):
            if h in alias_list:
                mapping[field] = i
                break

    if 'label' not in mapping and 'field_key' not in mapping:
        return None

    return mapping


def _safe_get(row_data, idx):
    """安全取值並 strip"""
    if idx is None or idx >= len(row_data):
        return ''
    val = row_data[idx]
    return str(val).strip() if val is not None else ''


def _row_to_spec_field(row_data, mapping, sort_order):
    """將一列資料轉為 SPEC field dict"""
    def get(key):
        return _safe_get(row_data, mapping.get(key))

    label = get('label')
    field_key = get('field_key')

    if not label and not field_key:
        return None

    type_str = get('type')
    formio_type = _resolve_formio_type(type_str)
    pg_type = _resolve_pg_type(get('pg_type'), formio_type)
    is_pii = _is_yes(get('pii'))
    required = _is_yes(get('required'))
    description = get('description')

    return {
        'field_key': field_key,
        'label': label,
        'formio_type': formio_type,
        'pg_type': pg_type,
        'constraints': {
            'required': required,
            'maxLength': None,
            'minLength': None,
            'min': None,
            'max': None,
            'pattern': None,
            'customValidation': None,
        },
        'is_pii': is_pii,
        'description': description,
        'default_value': None,
        'options': None,
        'grid_children': None,
        'lookup_category_code': None,
        'sort_order': sort_order,
    }


def _resolve_formio_type(type_str):
    """
    解析 Type 欄位值，回傳 formio_type 字串。
    接受中文標籤或英文值。
    """
    if not type_str:
        return 'textfield'

    type_str = type_str.strip()

    # 中文標籤
    if type_str in _LABEL_TO_FORMIO_TYPE:
        return _LABEL_TO_FORMIO_TYPE[type_str]

    # 英文值（不區分大小寫）
    normalized = _FORMIO_TYPE_NORMALIZE.get(type_str.lower())
    if normalized:
        return normalized

    logger.warning(f'未知的 Type 值 "{type_str}"，fallback 為 textfield')
    return 'textfield'


def _resolve_pg_type(pg_type_str, formio_type):
    """如果有提供 PG Type 就用，否則從 formio_type 推導"""
    if pg_type_str:
        return pg_type_str
    return _PG_DEFAULTS.get(formio_type, 'TEXT')


def _is_yes(val):
    """判斷 Y/是/true/1 等表示 True"""
    if not val:
        return False
    return val.upper() in ('Y', 'YES', '是', 'TRUE', '1')
